"""Core logic for the FX Deal ID Cleanup tool.

Pure functions only (no Streamlit imports) so they can be tested on their own.
Behaviour matches the original fx_cleanup.ipynb notebook.
"""
from __future__ import annotations

import io
import os
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl

EPS = 1e-6  # values smaller than this (absolute) are treated as zero
DEFAULT_DEAL_IN_ID = 627080  # main internal deal the residual exposure flows into
SOURCE_SHEET = "Sheet2"
TEMPLATE_SHEET = "Sheet3"
DATE_FMT = "DD-MMM-YYYY"


class CleanupError(Exception):
    """A problem with the uploaded workbook that the user can fix."""


@dataclass
class CleanupResult:
    output_bytes: bytes
    output_filename: str
    rows: list  # [(deal_out_id, trade_date, conv_date, quantity), ...]
    skipped_text_cells: list  # [(deal_id, month_label, raw_text), ...]
    first_row: int | None
    last_row: int | None
    deal_in_id: int
    warnings: list = field(default_factory=list)


# ---------------------------------------------------------------- helpers
def adjusted_10th(month_date: datetime) -> datetime:
    """10th of month_date's month, moved back to the nearest earlier weekday
    if the 10th is a Saturday or Sunday."""
    d = month_date.replace(day=10)
    while d.weekday() >= 5:  # Saturday=5, Sunday=6
        d -= timedelta(days=1)
    return d


def find_sheet2_header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row)):
        if row[0].value and str(row[0].value).strip().lower() == "risk month":
            return row[0].row
    raise CleanupError("Could not find the 'Risk Month' header row on Sheet2.")


def find_row_total_col(ws, header_row):
    for cell in ws[header_row]:
        if cell.value and "row total" in str(cell.value).strip().lower():
            return cell.column
    raise CleanupError("Could not find the 'Row Total (EUR)' column on Sheet2.")


def get_deal_id_columns(ws, header_row, first_deal_col):
    deal_ids = {}
    for c in range(first_deal_col, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, (int, float)):
            deal_ids[c] = int(v)
    return deal_ids


def get_month_rows(ws, header_row):
    rows = []
    r = header_row + 1
    while r <= ws.max_row:
        label = ws.cell(row=r, column=1).value
        date_val = ws.cell(row=r, column=2).value
        if label is None:
            r += 1
            continue
        if str(label).strip().lower() == "grand total":
            break
        if isinstance(date_val, datetime):
            rows.append((r, date_val))
        r += 1
    return rows


def extract_exposure_data(ws2):
    """Return (results, skipped_text_cells).

    results: [(deal_id, month_date, value)] for every nonzero numeric cell.
    skipped_text_cells: [(deal_id, month_label, raw_text)] for text cells
    (other than blank / '-').
    """
    header_row = find_sheet2_header_row(ws2)
    row_total_col = find_row_total_col(ws2, header_row)
    deal_ids = get_deal_id_columns(ws2, header_row, row_total_col + 1)
    month_rows = get_month_rows(ws2, header_row)

    results, skipped = [], []
    for r, month_date in month_rows:
        month_label = ws2.cell(row=r, column=1).value
        for c, deal_id in deal_ids.items():
            val = ws2.cell(row=r, column=c).value
            if val is None:
                continue
            if isinstance(val, str):
                if val.strip() not in ("-", ""):
                    skipped.append((deal_id, month_label, val))
                continue
            if abs(val) < EPS:
                continue
            results.append((deal_id, month_date, val))
    return results, skipped


def find_sheet3_header_row_and_columns(ws3):
    field_matchers = {
        "unique_id": "unique id",
        "deal_in_id": "deal in id",
        "deal_out_id": "deal out id",
        "trade_date": "trade date",
        "conv_date": "conv date",
        "quantity": "quantity",
    }
    for row in ws3.iter_rows(min_row=1, max_row=min(10, ws3.max_row)):
        headers = {c.column: str(c.value).strip().lower() for c in row if c.value}
        if any("deal in id" in h for h in headers.values()) and any(
            "deal out id" in h for h in headers.values()
        ):
            columns = {
                f: next((col for col, h in headers.items() if needle in h), None)
                for f, needle in field_matchers.items()
            }
            return row[0].row, columns
    raise CleanupError("Could not find the 'Deal In ID' / 'Deal Out ID' header row on Sheet3.")


def clear_previous_rows(ws3, start_row, columns):
    """Clear leftover example/previous-output values in the fields we write to.
    The Unique ID column is never touched."""
    write_cols = [c for f, c in columns.items() if f != "unique_id" and c is not None]
    if not write_cols:
        return
    last_data_row = start_row - 1
    r, empty_streak = start_row, 0
    while r <= ws3.max_row and empty_streak < 50:
        if any(ws3.cell(row=r, column=c).value is not None for c in write_cols):
            last_data_row, empty_streak = r, 0
        else:
            empty_streak += 1
        r += 1
    for r in range(start_row, last_data_row + 1):
        for c in write_cols:
            cell = ws3.cell(row=r, column=c)
            cell.value = None
            cell.number_format = "General"


def write_results(ws3, header_row, columns, results, trade_date, deal_in_id):
    if columns.get("deal_in_id") is None or columns.get("deal_out_id") is None:
        raise CleanupError("Sheet3 is missing the Deal In ID / Deal Out ID columns.")

    start_row = header_row + 3  # header, note row, blank row, then data
    clear_previous_rows(ws3, start_row, columns)

    out_rows = []
    for i, (deal_id, month_date, qty) in enumerate(results):
        r = start_row + i
        conv_date = adjusted_10th(month_date)

        ws3.cell(row=r, column=columns["deal_in_id"], value=deal_in_id).number_format = "General"
        ws3.cell(row=r, column=columns["deal_out_id"], value=deal_id).number_format = "General"
        if columns.get("trade_date") is not None:
            ws3.cell(row=r, column=columns["trade_date"], value=trade_date).number_format = DATE_FMT
        if columns.get("conv_date") is not None:
            ws3.cell(row=r, column=columns["conv_date"], value=conv_date).number_format = DATE_FMT
        if columns.get("quantity") is not None:
            ws3.cell(row=r, column=columns["quantity"], value=qty).number_format = "#,##0.00"
        out_rows.append((deal_id, trade_date, conv_date, qty))

    last_row = start_row + len(results) - 1 if results else start_row - 1
    return start_row, last_row, out_rows


# ---------------------------------------------------------------- entry point
def process_workbook(
    file_bytes: bytes,
    filename: str,
    trade_date: date | None = None,
    deal_in_id: int = DEFAULT_DEAL_IN_ID,
) -> CleanupResult:
    """Run the full cleanup on an uploaded workbook and return the new file."""
    base, ext = os.path.splitext(filename)
    ext = ext.lower()
    if ext not in (".xlsx", ".xlsm"):
        raise CleanupError("Please upload an .xlsx or .xlsm file (old .xls files are not supported).")

    trade_dt = datetime.combine(trade_date or date.today(), datetime.min.time())

    try:
        wb_data = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        wb_out = openpyxl.load_workbook(
            io.BytesIO(file_bytes), data_only=False, keep_vba=(ext == ".xlsm")
        )
    except Exception as exc:  # corrupt / not really an Excel file
        raise CleanupError(f"Could not open the workbook: {exc}") from exc

    if SOURCE_SHEET not in wb_data.sheetnames or TEMPLATE_SHEET not in wb_out.sheetnames:
        raise CleanupError(
            f"Expected sheets named '{SOURCE_SHEET}' (source data) and "
            f"'{TEMPLATE_SHEET}' (output template). Found: {', '.join(wb_data.sheetnames)}."
        )

    results, skipped = extract_exposure_data(wb_data[SOURCE_SHEET])

    ws3 = wb_out[TEMPLATE_SHEET]
    header_row, columns = find_sheet3_header_row_and_columns(ws3)
    first_row, last_row, out_rows = write_results(
        ws3, header_row, columns, results, trade_dt, deal_in_id
    )

    buf = io.BytesIO()
    wb_out.save(buf)

    return CleanupResult(
        output_bytes=buf.getvalue(),
        output_filename=f"{base}_output{ext}",
        rows=out_rows,
        skipped_text_cells=skipped,
        first_row=first_row if results else None,
        last_row=last_row if results else None,
        deal_in_id=deal_in_id,
    )
